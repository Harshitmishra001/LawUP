import io
import docx
import pypdf
import pytest
from fastapi.testclient import TestClient

from backend.app.agents.ingestion import IngestionAgent, ClauseUnit
from backend.app.api.main import app


@pytest.fixture
def agent():
    return IngestionAgent(min_clause_words=5)


@pytest.fixture
def client():
    return TestClient(app)


def test_ingest_plain_text(agent):
    text = """
    MASTER SERVICES AGREEMENT

    This Master Services Agreement is dated October 1, 2026.

    1. SCOPE OF SERVICES
    Vendor shall provide cloud infrastructure monitoring and maintenance services to Customer.

    2. FEES AND INVOICING
    Customer agrees to pay all undisputed invoices within forty-five calendar days of receipt.

    3. LIMITATION OF LIABILITY
    Neither party shall be liable for indirect, incidental, special, or consequential damages.
    """
    units = agent.ingest(text)
    assert len(units) >= 3
    assert all(isinstance(u, ClauseUnit) for u in units)
    assert any("SCOPE OF SERVICES" in u.heading for u in units)
    assert any("FEES AND INVOICING" in u.heading for u in units)


def test_backward_compatibility_run(agent):
    text = """
    1. TERMINATION
    Client may terminate this contract upon thirty days written notice.

    2. GOVERNING LAW
    This agreement shall be governed by the laws of the State of Delaware.
    """
    clauses = agent.run(text)
    assert isinstance(clauses, list)
    assert len(clauses) >= 2
    assert all(isinstance(c, str) for c in clauses)


def test_ingest_docx(agent):
    doc = docx.Document()
    doc.add_paragraph("ARTICLE 1: NON-SOLICITATION\nEmployee shall not solicit any customers of the Employer for a period of one year.")
    doc.add_paragraph("ARTICLE 2: SEVERABILITY\nIf any provision of this agreement is held invalid, the remainder shall remain in effect.")

    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)

    units = agent.ingest(stream.getvalue(), filename="agreement.docx")
    assert len(units) == 2
    assert "ARTICLE 1: NON-SOLICITATION" in units[0].heading
    assert "Employee shall not solicit" in units[0].text


def test_ingest_pdf(agent):
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=72, height=72)
    stream = io.BytesIO()
    writer.write(stream)
    stream.seek(0)

    units = agent.ingest(stream.getvalue(), filename="empty.pdf")
    assert isinstance(units, list)


def test_minimum_words_filtering(agent):
    short_text = """
    1. TITLE
    Too short.

    2. SUBSTANTIVE SECTION
    This section contains enough words to satisfy the substantive clause length threshold requirement.
    """
    units = agent.ingest(short_text)
    assert len(units) == 1
    assert "SUBSTANTIVE SECTION" in units[0].heading


def test_api_ingest_endpoint(client):
    contract_text = """
    1. SERVICES
    Contractor shall perform backend architecture design and development for Company.

    2. PAYMENT
    Company shall pay Contractor within thirty days of approved milestone completion.
    """
    response = client.post("/ingest", json={"contract_text": contract_text})
    assert response.status_code == 200
    data = response.json()
    assert data["total_clauses"] == 2
    assert len(data["clauses"]) == 2
    assert "1. SERVICES" in data["clauses"][0]["heading"]
    assert "2. PAYMENT" in data["clauses"][1]["heading"]


def test_api_ingest_file_endpoint(client):
    doc = docx.Document()
    doc.add_paragraph("1. TERM\nThis agreement shall remain in effect for two calendar years.")
    doc.add_paragraph("2. JURISDICTION\nAny legal proceedings must be initiated in courts in California.")

    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)

    files = {"file": ("test_contract.docx", stream.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    response = client.post("/ingest/file", files=files)
    assert response.status_code == 200
    data = response.json()
    assert data["total_clauses"] == 2
    assert "1. TERM" in data["clauses"][0]["heading"]
